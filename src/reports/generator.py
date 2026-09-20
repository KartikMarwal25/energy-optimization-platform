"""Accessible, plain-language exports for GridPulse."""
import base64
import html
import os
from datetime import datetime
import pandas as pd

REPORT_DIR = "reports"
NAVY, TEAL, PALE, SLATE = "102A43", "087E8B", "EAF4F4", "627D98"

def _mkdir(): os.makedirs(REPORT_DIR, exist_ok=True)
def _ctx(context): return context or {}

def export_csv(df: pd.DataFrame, name: str):
    _mkdir(); path = os.path.join(REPORT_DIR, f"{name}.csv"); df.to_csv(path, index=False); return path

def export_excel(df: pd.DataFrame, name: str):
    _mkdir(); path = os.path.join(REPORT_DIR, f"{name}.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        ws = writer.sheets["Data"]; ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            from copy import copy
            cell.font = copy(cell.font); cell.font = cell.font.copy(bold=True, color="FFFFFF")
            cell.fill = copy(cell.fill); cell.fill = cell.fill.copy(fgColor=NAVY, fill_type="solid")
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in cells[:101]) + 2, 12), 38)
    return path

def _cards(context):
    return "".join(f"<div class='kpi'><span>{html.escape(str(x.get('label','Metric')))}</span><strong>{html.escape(str(x.get('value','—')))}</strong><small>{html.escape(str(x.get('detail','')))}</small></div>" for x in context.get("kpis", []))

def _charts(paths):
    output = ""
    for path in paths or []:
        if os.path.exists(path):
            with open(path, "rb") as f: encoded = base64.b64encode(f.read()).decode("ascii")
            caption = os.path.splitext(os.path.basename(path))[0].replace("_", " ").title()
            output += f"<figure><img src='data:image/png;base64,{encoded}' alt='{html.escape(caption)}'><figcaption>{html.escape(caption)}</figcaption></figure>"
    return output

def export_html(summary_text, tables=None, charts=None, name="summary", report_context=None):
    _mkdir(); c = _ctx(report_context); path = os.path.join(REPORT_DIR, f"{name}.html")
    findings = "".join(f"<li><strong>{html.escape(str(x.get('title','Finding')))}</strong><br>{html.escape(str(x.get('detail','')))}</li>" for x in c.get("findings", [])) or "<li>No decision-ready findings are available. Run the relevant analysis before exporting.</li>"
    method = "".join(f"<li>{html.escape(str(x))}</li>" for x in c.get("methodology", []))
    table_html = "".join(f"<section><h2>{html.escape(title)}</h2>{df.head(20).to_html(index=False, classes='report-table', border=0)}</section>" for title, df in (tables or {}).items() if df is not None and not df.empty)
    summary = "<br>".join(html.escape(x) for x in summary_text.splitlines() if x.strip())
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"""<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(c.get('title','Energy Operations Report'))}</title><style>
        body{{font:15px Arial,sans-serif;line-height:1.5;color:#{NAVY};background:#f5f8fa;margin:0}}main{{max-width:980px;margin:auto;background:#fff;padding:46px 60px}}header{{border-bottom:4px solid #{TEAL};padding-bottom:20px}}h1{{margin:0;font-size:30px}}h2{{font-size:19px;margin-top:32px}}.muted,small,figcaption{{color:#{SLATE}}}.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin:25px 0}}.kpi{{background:#{PALE};border-left:4px solid #{TEAL};padding:14px}}.kpi span,.kpi small{{display:block;font-size:12px}}.kpi strong{{font-size:23px;display:block}}.report-table{{border-collapse:collapse;width:100%;font-size:12px}}th{{background:#{NAVY};color:#fff;text-align:left;padding:8px}}td{{padding:8px;border:1px solid #d9e2ec;vertical-align:top}}tr:nth-child(even){{background:#f5f8fa}}figure{{margin:28px 0}}figure img{{max-width:100%;border:1px solid #d9e2ec}}footer{{margin-top:36px;border-top:1px solid #d9e2ec;padding-top:14px;color:#{SLATE};font-size:11px}}@media print{{body{{background:#fff}}main{{padding:20px}}}}</style></head><body><main><header><h1>{html.escape(c.get('title','Energy Operations Report'))}</h1><div class='muted'>{html.escape(c.get('subtitle','Decision support for smart-grid operations'))}</div><div class='muted'>Generated {html.escape(c.get('generated_at',datetime.now().strftime('%d %B %Y, %H:%M')))}</div></header><div class='kpis'>{_cards(c)}</div><section><h2>Executive summary</h2><p>{summary}</p></section><section><h2>Findings and next steps</h2><ol>{findings}</ol></section>{_charts(charts)}<section><h2>How to read this report</h2><ul>{method}</ul></section>{table_html}<footer>GridPulse decision support. Validate forecasts and anomaly flags against operational records before taking action.</footer></main></body></html>""")
    return path

def export_pdf(summary_text, tables=None, charts=None, name="summary", report_context=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    _mkdir(); c = _ctx(report_context); path = os.path.join(REPORT_DIR, f"{name}.pdf")
    styles = getSampleStyleSheet(); styles.add(ParagraphStyle("ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=29, textColor=colors.HexColor('#'+NAVY), spaceAfter=5)); styles.add(ParagraphStyle("Section", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor('#'+NAVY), spaceBefore=18, spaceAfter=8)); styles.add(ParagraphStyle("TableCell", parent=styles["BodyText"], fontSize=7, leading=9)); styles.add(ParagraphStyle("TableHead", parent=styles["TableCell"], textColor=colors.white, fontName="Helvetica-Bold"))
    story=[Paragraph(html.escape(c.get('title','Energy Operations Report')),styles['ReportTitle']),Paragraph(html.escape(c.get('subtitle','Decision support for smart-grid operations')),styles['BodyText']),Spacer(1,12)]
    if c.get('kpis'):
        row=[Paragraph(f"<b>{html.escape(str(x.get('label','Metric')))}</b><br/><font size='15'>{html.escape(str(x.get('value','—')))}</font><br/><font color='#{SLATE}'>{html.escape(str(x.get('detail','')))}</font>",styles['BodyText']) for x in c['kpis'][:4]]
        t=Table([row],colWidths=[(A4[0]-84)/len(row)]*len(row));t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#'+PALE)),('BOX',(0,0),(-1,-1),.5,colors.HexColor('#'+TEAL)),('INNERGRID',(0,0),(-1,-1),.5,colors.white),('VALIGN',(0,0),(-1,-1),'TOP'),('PADDING',(0,0),(-1,-1),10)]));story += [t,Spacer(1,8)]
    story += [Paragraph('Executive summary',styles['Section']),Paragraph('<br/>'.join(html.escape(x) for x in summary_text.splitlines() if x.strip()),styles['BodyText']),Paragraph('Findings and next steps',styles['Section'])]
    for x in c.get('findings',[])[:5]: story += [Paragraph(f"<b>{html.escape(str(x.get('title','Finding')))}</b><br/>{html.escape(str(x.get('detail','')))}",styles['BodyText']),Spacer(1,5)]
    for chart in charts or []:
        if os.path.exists(chart): story += [Paragraph(os.path.splitext(os.path.basename(chart))[0].replace('_',' ').title(),styles['Section']),Image(chart,width=6.8*inch,height=3.7*inch)]
    story.append(Paragraph('How to read this report',styles['Section']))
    for x in c.get('methodology',[]): story.append(Paragraph('• '+html.escape(str(x)),styles['BodyText']))
    for title,df in (tables or {}).items():
        if df is None or df.empty: continue
        story.append(Paragraph(html.escape(title),styles['Section'])); clipped=df.head(12).iloc[:,:6]; rows=[[Paragraph(html.escape(str(v)),styles['TableHead']) for v in clipped.columns]]+[[Paragraph(html.escape(str(v))[:120],styles['TableCell']) for v in row] for _,row in clipped.iterrows()];t=Table(rows,repeatRows=1);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+NAVY)),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9E2EC')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FA')]),('PADDING',(0,0),(-1,-1),5)]));story.append(t)
    SimpleDocTemplate(path,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=45,bottomMargin=40).build(story); return path

def export_pptx(summary_text, charts=None, dataframes=None, name="summary", report_context=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    _mkdir(); c=_ctx(report_context); path=os.path.join(REPORT_DIR,f"{name}.pptx"); prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);navy=RGBColor(16,42,67);teal=RGBColor(8,126,139);slate=RGBColor(98,125,152)
    def title(slide,text,sub=''):
        box=slide.shapes.add_textbox(Inches(.7),Inches(.35),Inches(12),Inches(.8));p=box.text_frame.paragraphs[0];p.text=text;p.font.size=Pt(29);p.font.bold=True;p.font.color.rgb=navy
        if sub: box=slide.shapes.add_textbox(Inches(.72),Inches(1.08),Inches(11.5),Inches(.35));p=box.text_frame.paragraphs[0];p.text=sub;p.font.size=Pt(12);p.font.color.rgb=slate
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=navy;box=slide.shapes.add_textbox(Inches(.8),Inches(2.3),Inches(11.5),Inches(1.1));p=box.text_frame.paragraphs[0];p.text=c.get('title','Energy Operations Report');p.font.size=Pt(35);p.font.bold=True;p.font.color.rgb=RGBColor(255,255,255);box=slide.shapes.add_textbox(Inches(.82),Inches(3.5),Inches(10.5),Inches(.6));p=box.text_frame.paragraphs[0];p.text=c.get('subtitle','Decision support for smart-grid operations');p.font.size=Pt(17);p.font.color.rgb=RGBColor(220,240,243)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);title(slide,'Executive overview','Validate findings against operational records before acting.')
    for i,x in enumerate(c.get('kpis',[])[:4]):
        box=slide.shapes.add_textbox(Inches(.7+i*3.15),Inches(1.7),Inches(2.8),Inches(1.4));box.fill.solid();box.fill.fore_color.rgb=RGBColor(234,244,244);tf=box.text_frame;tf.margin_left=Inches(.15);tf.margin_top=Inches(.12);p=tf.paragraphs[0];p.text=str(x.get('label','Metric'));p.font.size=Pt(11);p.font.color.rgb=slate;p=tf.add_paragraph();p.text=str(x.get('value','—'));p.font.size=Pt(24);p.font.bold=True;p.font.color.rgb=teal;p=tf.add_paragraph();p.text=str(x.get('detail',''));p.font.size=Pt(9);p.font.color.rgb=slate
    box=slide.shapes.add_textbox(Inches(.8),Inches(3.65),Inches(11.5),Inches(2.2));tf=box.text_frame;tf.word_wrap=True
    for i,x in enumerate(c.get('findings',[])[:4]): p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=f"{x.get('title','Finding')}: {x.get('detail','')}";p.font.size=Pt(16);p.font.color.rgb=navy;p.space_after=Pt(10)
    for chart in charts or []:
        if os.path.exists(chart): slide=prs.slides.add_slide(prs.slide_layouts[6]);title(slide,os.path.splitext(os.path.basename(chart))[0].replace('_',' ').title());slide.shapes.add_picture(chart,Inches(.75),Inches(1.35),width=Inches(11.8))
    for heading,df in (dataframes or {}).items():
        if df is None or df.empty: continue
        slide=prs.slides.add_slide(prs.slide_layouts[6]);title(slide,heading,'Top records shown for discussion');clip=df.head(8).iloc[:,:5];table=slide.shapes.add_table(len(clip)+1,len(clip.columns),Inches(.65),Inches(1.45),Inches(12),Inches(4.9)).table
        for j,col in enumerate(clip.columns): cell=table.cell(0,j);cell.text=str(col);cell.fill.solid();cell.fill.fore_color.rgb=navy
        for i,(_,row) in enumerate(clip.iterrows(),1):
            for j,col in enumerate(clip.columns): table.cell(i,j).text=str(row[col])[:80]
    prs.save(path);return path
